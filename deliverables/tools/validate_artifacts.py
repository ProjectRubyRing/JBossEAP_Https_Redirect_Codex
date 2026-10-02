"""Validate local report completeness, Excel formatting, and script text."""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from build_report import BOOK, MD, ROOT, FONT_NAME, SHEET_NAMES, REF_DEF, blocks, plain


def validate():
    results = []

    def check(name, condition, detail=""):
        results.append({"check": name, "result": "PASS" if condition else "FAIL", "detail": detail})
        if not condition:
            raise AssertionError(f"{name}: {detail}")

    with zipfile.ZipFile(BOOK) as z:
        check("XLSX ZIP integrity", z.testzip() is None)
    wb = load_workbook(BOOK)
    check("20 chapter worksheets", wb.sheetnames == SHEET_NAMES, str(len(wb.sheetnames)))
    fonts = sorted({f.name for f in wb._fonts})
    check("All workbook fonts including comments are Meiryo UI", fonts == [FONT_NAME], str(fonts))

    colors = []
    for font in wb._fonts:
        if font.color and font.color.type == "rgb":
            colors.append(font.color.rgb)
    for fill in wb._fills:
        for color in (fill.fgColor, fill.bgColor):
            if color.type == "rgb":
                colors.append(color.rgb)
    for border in wb._borders:
        for side_name in ("left", "right", "top", "bottom"):
            side = getattr(border, side_name)
            if side and side.color and side.color.type == "rgb":
                colors.append(side.color.rgb)
    check("Explicit style colors are grayscale", all(c[-6:-4] == c[-4:-2] == c[-2:] for c in colors), str(sorted(set(colors))))
    all_cells = [c for ws in wb for cells in ws for c in cells if c.value is not None]
    values = {str(c.value) for c in all_cells}
    check("Literal code / no formulas", not any(c.data_type == "f" for c in all_cells))
    check("No Unicode replacement characters", not any("\ufffd" in str(c.value) for c in all_cells))
    check("All nonempty cells have Meiryo UI", all(c.font.name == FONT_NAME for c in all_cells))
    check("All nonempty cells wrap", all(c.alignment.wrap_text for c in all_cells))
    check("All chapters freeze top four rows", all(ws.freeze_panes == "A5" for ws in wb))
    check("All chapters use landscape A4", all(ws.page_setup.orientation == "landscape" and int(ws.page_setup.paperSize) == 9 for ws in wb))
    check("All chapters define print areas", all(bool(ws.print_area) for ws in wb))
    check("All chapters fit width, allow multiple pages vertically", all(ws.sheet_properties.pageSetUpPr.fitToPage and (ws.page_setup.fitToWidth if ws.page_setup.fitToWidth is not None else 1) == 1 and ws.page_setup.fitToHeight == 0 for ws in wb))
    heights = [d.height for ws in wb for d in ws.row_dimensions.values() if d.height is not None]
    check("Explicit row heights within Excel limit", all(0 < h <= 409 for h in heights), f"maximum={max(heights)}pt")
    check("Diagram embedded in principles chapter", len(wb["03_原理と通信図"]._images) == 1)
    check("Diagram file exists for Markdown", ROOT.joinpath("assets/https_flow.png").is_file())

    text = MD.read_text(encoding="utf-8")
    refs = {m.group(1): m.group(2) for m in REF_DEF.finditer(text)}
    check("30 primary source definitions", len(refs) == 30)
    referenced = re.findall(r"\[[^\]]+\]\[(S\d+)\]", text)
    check("No undefined source references", all(key in refs for key in referenced))
    external_links = {
        c.hyperlink.target + ("#" + c.hyperlink.location if c.hyperlink.location else "")
        for c in all_cells if c.hyperlink and c.hyperlink.target
    }
    check("Source URLs available as native Excel hyperlinks", set(refs.values()).issubset(external_links))
    body = REF_DEF.sub("", text)
    chunks = re.split(r"^## (.+)$", body, flags=re.M)
    expected = []
    for chunk in [chunks[0]] + [chunks[i] for i in range(2, len(chunks), 2)]:
        for block in blocks(chunk.splitlines()):
            if block["type"] in ("paragraph", "item", "heading") and not block.get("text", "").startswith("# "):
                expected.append(plain(block["text"]))
            elif block["type"] == "table":
                expected.extend(plain(v) for row in [block["headers"]] + block["rows"] for v in row)
            elif block["type"] == "code":
                expected.extend(line for line in block["lines"] if line)
    missing = [v for v in expected if v not in values]
    check("Markdown paragraphs, tables and code fully represented in Excel", not missing, f"checked={len(expected)} items, missing={len(missing)}")

    text_paths = list(ROOT.joinpath("scripts").glob("*")) + list(ROOT.joinpath("tools").glob("*.py")) + [MD]
    check("Scripts and Markdown are valid UTF-8", all(p.read_bytes().decode("utf-8") is not None for p in text_paths))
    check("CLI and Linux scripts use LF without BOM", all(not p.read_bytes().startswith(b"\xef\xbb\xbf") and b"\r\n" not in p.read_bytes() for p in ROOT.joinpath("scripts").glob("*") if p.suffix in (".cli", ".sh")))
    whitespace_errors = []
    for path in text_paths:
        for number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            if line.rstrip() != line:
                whitespace_errors.append(f"{path.name}:{number}")
    check("No unintended trailing whitespace", not whitespace_errors, str(whitespace_errors))
    check("PowerShell path comparison respects case-sensitive URI paths", "AbsolutePath -cne" in ROOT.joinpath("scripts/Test-HttpsUploadRedirect.ps1").read_text(encoding="utf-8"))

    out = ROOT / "artifact_validation.json"
    out.write_text(json.dumps({"results": results, "scope": "Local deliverables only; no EAP or ALB runtime test"}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"PASS: {len(results)} artifact checks; {len(all_cells)} nonempty cells; {len(expected)} source content items.")


if __name__ == "__main__":
    validate()
