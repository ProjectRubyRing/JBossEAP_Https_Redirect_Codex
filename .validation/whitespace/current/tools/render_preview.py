"""Render real workbook cells for layout inspection; this is NOT Excel rendering."""
from __future__ import annotations

import json
import math
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw

from build_report import BOOK, ROOT, COL_PIXELS, ui_font, wrap

PREVIEW = ROOT / "assets" / "previews"


def rgb(color, fallback):
    if color is not None and color.type == "rgb":
        return "#" + color.rgb[-6:]
    return fallback


def render():
    PREVIEW.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(BOOK)
    records = []
    for ws in wb:
        x = [0]
        for width in COL_PIXELS:
            x.append(x[-1] + width)
        y = [0]
        for row in range(1, ws.max_row + 1):
            points = ws.row_dimensions[row].height or ws.sheet_format.defaultRowHeight or 23
            y.append(y[-1] + math.ceil(points * 96 / 72))
        img = Image.new("RGB", (x[-1], y[-1] + 42), "white")
        d = ImageDraw.Draw(img)
        merged = {(m.min_row, m.min_col): m for m in ws.merged_cells.ranges}
        skip = set()
        for m in ws.merged_cells.ranges:
            for row in range(m.min_row, m.max_row + 1):
                for col in range(m.min_col, m.max_col + 1):
                    if (row, col) != (m.min_row, m.min_col):
                        skip.add((row, col))
        overflows = []
        for row in range(1, ws.max_row + 1):
            for col in range(1, 7):
                if (row, col) in skip:
                    continue
                cell = ws.cell(row, col)
                region = merged.get((row, col))
                end_col = region.max_col if region else col
                end_row = region.max_row if region else row
                box = (x[col - 1], y[row - 1], x[end_col], y[end_row])
                fill = rgb(cell.fill.fgColor, "#FFFFFF") if cell.fill.patternType == "solid" else "#FFFFFF"
                d.rectangle(box, fill=fill)
                if cell.border.left.style or cell.border.right.style or cell.border.top.style or cell.border.bottom.style:
                    d.rectangle((box[0], box[1], box[2] - 1, box[3] - 1), outline="#CCCCCC", width=1)
                if cell.value is None:
                    continue
                size = float(cell.font.sz or 11)
                font = ui_font(size, bool(cell.font.b))
                line_height = round(size * 1.55 * 96 / 72)
                lines = wrap(str(cell.value), font, box[2] - box[0] - 24)
                yy = box[1] + 5
                for line in lines:
                    d.text((box[0] + 12, yy), line, font=font, fill=rgb(cell.font.color, "#222222"))
                    yy += line_height
                if yy - line_height + font.size + 5 > box[3]:
                    overflows.append(cell.coordinate)
        for picture in ws._images:
            # The generated workbook uses a two-cell image anchor.
            anchor = picture.anchor
            start = anchor._from
            end = anchor.to
            x1 = x[start.col] + round(start.colOff / 9525)
            y1 = y[start.row] + round(start.rowOff / 9525)
            x2 = x[end.col] + round(end.colOff / 9525)
            y2 = y[end.row] + round(end.rowOff / 9525)
            with Image.open(ROOT / "assets" / "https_flow.png") as original:
                resized = original.resize((x2 - x1, y2 - y1), Image.Resampling.LANCZOS)
                img.paste(resized, (x1, y1))
        d.text((12, y[-1] + 9), "Meiryo UI / セルデータからの Pillow 描画プレビュー（Excel 画面ではありません）", font=ui_font(10), fill="#555555")
        img.save(PREVIEW / f"{ws.title}.png", optimize=True)
        # A readable first-screen crop supports quick inspection.
        img.crop((0, 0, img.width, min(1600, img.height))).save(PREVIEW / f"{ws.title}_excerpt.png", optimize=True)
        records.append({"sheet": ws.title, "width": img.width, "height": img.height, "detected_text_overflows": overflows})
    result = {"status": "PASS" if not any(r["detected_text_overflows"] for r in records) else "FAIL", "renderer": "Pillow with installed Meiryo UI; not native Excel", "records": records}
    ROOT.joinpath("preview_validation.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if result["status"] != "PASS":
        raise AssertionError(json.dumps(result, ensure_ascii=False))
    print(f"PASS: {len(records)} workbook data previews; no detected text overflow.")


if __name__ == "__main__":
    render()
